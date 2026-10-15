#!/usr/bin/env python3
"""
ATL Activity Generator (server version)
=======================================
Upload a class text (grades 9-12) and get a ready-to-print ATL activity as a Word (.docx) or PDF file.

Two versions of the activity are produced from the same generated content:
  * "teacher"  - student pages + running-record copy + vocabulary targets, rubrics, full-credit
                 answers, writing observations and a score report
  * "student"  - student-facing pages only (reading, vocabulary, comprehension questions,
                 written responses) with no answers, plus the teacher running-record copy as the
                 last page, marked to remove before handing the packet out

Command line:
    export ANTHROPIC_API_KEY=sk-ant-...
    python atl_generator.py reading.pdf --title "The Balance of Soft Power" --grade 11 \
        --subject "Global Politics" --decode-words 410 --mode teacher --format pdf --out soft_power.pdf

Web app:  streamlit run app.py
"""
from __future__ import annotations

import argparse
import io
import json
import os
import re
import sys
from pathlib import Path

DEFAULT_MODEL = os.environ.get("ATL_MODEL", "claude-sonnet-5-5")
NAVY, RED, GRAY = "1B234B", "C00000", "808080"

# ======================================================================================
# 1. Reading the uploaded text
# ======================================================================================


def _pdf_text(data: bytes) -> str:
    """Extract text from a PDF, rebuilding paragraphs from vertical gaps between lines."""
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    frags = []  # [y, size, text, page]

    for pi, page in enumerate(reader.pages):
        cur = {"line": None}

        def visitor(text, cm, tm, font_dict, font_size, _pi=pi, _cur=cur):
            if not text or not text.strip("\n"):
                return
            try:
                y = tm[4] * cm[2] + tm[5] * cm[3] + cm[5]
                size = abs(font_size * tm[3] * cm[3]) or 10
                if size < 4:
                    size = abs(tm[3] * cm[3]) or 10
            except Exception:
                y, size = 0, 10
            line = _cur["line"]
            if line and abs(line[0] - y) < size * 0.5:
                line[2] += text.replace("\n", " ")
            else:
                line = [y, size, text.replace("\n", " "), _pi]
                frags.append(line)
                _cur["line"] = line

        page.extract_text(visitor_text=visitor)

    paras, buf, prev = [], "", None
    for y, size, text, pg in frags:
        s = re.sub(r"\s+", " ", text).strip()
        if not s or re.fullmatch(r"\d{1,3}", s):  # blanks and bare page numbers
            continue
        new_para = prev and prev[3] == pg and (prev[0] - y) > prev[1] * 1.7
        if new_para:
            paras.append(buf)
            buf = s
        elif re.search(r"[A-Za-z]-$", buf) and re.match(r"[a-z]", s):  # re-join hyphenated words
            buf = buf[:-1] + s
        else:
            buf = f"{buf} {s}" if buf else s
        prev = (y, size, text, pg)
    if buf:
        paras.append(buf)
    return "\n\n".join(paras)


def extract_text(data: bytes, ext: str) -> str:
    """Extract plain text from .txt/.md/.docx/.pdf bytes."""
    ext = ext.lower()
    if ext in (".txt", ".md"):
        return data.decode("utf-8", errors="ignore")
    if ext == ".docx":
        from docx import Document

        d = Document(io.BytesIO(data))
        return "\n\n".join(p.text for p in d.paragraphs if p.text.strip())
    if ext == ".pdf":
        text = _pdf_text(data)
        if len(text.split()) < 30:
            raise ValueError("This PDF has little or no selectable text (it may be a scan). Paste the text instead.")
        return text
    raise ValueError(f"Unsupported file type: {ext} (use .pdf, .docx, .txt or .md)")


def read_file(path: str) -> str:
    p = Path(path)
    return extract_text(p.read_bytes(), p.suffix)


def to_paragraphs(text: str) -> list[str]:
    """Split on blank lines; join hard-wrapped lines; strip leading '(1)' paragraph numbers."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    paras = []
    for block in re.split(r"\n\s*\n", text):
        block = " ".join(line.strip() for line in block.split("\n") if line.strip())
        block = re.sub(r"^\(?\d{1,2}\)\s+", "", block)
        if block:
            paras.append(block)
    return paras


def word_count(s: str) -> int:
    return len(s.split())


def split_for_decoding(paras: list[str], decode_words: int):
    """Running record covers whole paragraphs until >= decode_words (0 = whole text)."""
    if not decode_words:
        return paras, []
    total, cut = 0, len(paras)
    for i, p in enumerate(paras):
        total += word_count(p)
        if total >= decode_words:
            cut = i + 1
            break
    return paras[:cut], paras[cut:]


# ======================================================================================
# 2. Generating the assessment content with Claude
# ======================================================================================

SYSTEM_PROMPT = """You are an expert secondary literacy specialist who writes diagnostic reading \
assessments (running record + vocabulary + comprehension + writing) for grades 9-12. \
You write precise, text-dependent items with clear scoring guidance. Respond with a single JSON \
object and nothing else (no markdown fences, no commentary)."""

USER_PROMPT = """Create the content for an ATL (Approaches to Learning) literacy skill assessment.

GRADE LEVEL: {grade}
SUBJECT/COURSE: {subject}
TEXT TITLE: {title}

SECTION A - the passage the student reads aloud (running record) and is questioned on:
<section_a>
{section_a}
</section_a>

SECTION B - additional text read silently, used only for the written responses (may be empty):
<section_b>
{section_b}
</section_b>

Return JSON with exactly this structure:
{{
  "short_title": "<=6 words for the file name and footer",
  "vocab": [  // EXACTLY 5, in order of appearance in Section A
    {{"sentence": "<verbatim excerpt from Section A, <=35 words, containing the target>",
      "target": "<the target word/phrase, EXACT substring of the sentence>",
      "teacher_target": "<concise student-friendly definition the teacher scores against>"}}
  ],
  "preview_words": ["<5 more challenging words from the whole text for informal, unscored checks>"],
  "proper_nouns": ["<proper nouns / technical terms the teacher should note separately and NOT count as decoding errors>"],
  "questions": [  // EXACTLY 5, in this order and with these labels
    {{"type": "basic explicit", "question": "...", "full_credit": "...", "distinction": ""}},
    {{"type": "complex explicit", "question": "...", "full_credit": "...", "distinction": "..."}},
    {{"type": "basic inferential", "question": "...", "full_credit": "...", "distinction": ""}},
    {{"type": "cause-and-effect inferential", "question": "...", "full_credit": "...", "distinction": "..."}},
    {{"type": "complex inferential", "question": "...", "full_credit": "...", "distinction": "..."}}
  ],
  "written": {{
    "summary_prompt": "Summarize the main idea of the text below.",
    "thesis_prompt": "<an arguable, open-ended thesis-practice question linking the text to a broader issue>",
    "comprehension_prompt": "<a question requiring analysis AND the instruction 'Use specific evidence from the text to explain your answer.'>"
  }},
  "lexile_estimate": "<approximate Lexile range for the text, e.g. 1050L-1150L; this is only an estimate>"
}}

RULES
- Vocabulary: choose Tier 2/Tier 3 words that are challenging for grade {grade} and whose meaning a student must \
actually know (not guessable as a trivial synonym); not proper nouns; spread across Section A; each sentence must be \
copied verbatim from Section A. Teacher targets are short and scoreable.
- Questions come only from Section A. Explicit = answer stated in text. Implicit/inferential = answer must be \
reasoned from the text. "full_credit" lists all key elements a full-credit (2-point) answer must include.
- For complex/cause-and-effect items, "distinction" contrasts a student who merely retrieves a fact (give a short \
sample quote) with a student whose answer shows deeper understanding (give a sample quote). Format: \
'Draw a distinction between the student who says: "..." and has ... and the student who says/explains ...'. \
Use "" when not needed.
- The thesis prompt and comprehension prompt may draw on the whole text (Sections A and B).
- Do not invent facts that are not in the text."""


def _parse_json(raw: str) -> dict:
    raw = raw.strip()
    raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw)
    return json.loads(raw[raw.find("{") : raw.rfind("}") + 1])


def _norm(s: str) -> str:
    s = s.replace("\u2019", "'").replace("\u2018", "'").replace("\u201c", '"').replace("\u201d", '"')
    return re.sub(r"\s+", " ", s).strip().lower()


def validate_content(content: dict, section_a: str) -> list[str]:
    """Warnings the teacher should look at (never raises on soft problems)."""
    warnings = []
    if len(content.get("vocab", [])) != 5:
        warnings.append(f"Expected 5 vocabulary items, got {len(content.get('vocab', []))}.")
    if len(content.get("questions", [])) != 5:
        warnings.append(f"Expected 5 comprehension questions, got {len(content.get('questions', []))}.")
    a = _norm(section_a)
    for i, v in enumerate(content.get("vocab", []), 1):
        if _norm(v.get("target", "")) not in _norm(v.get("sentence", "")):
            warnings.append(f"Vocab {i}: target '{v.get('target')}' is not in its sentence.")
        if _norm(v.get("sentence", "")).rstrip(".\u2026") not in a:
            warnings.append(f"Vocab {i}: sentence is not a verbatim excerpt of the text. Check it.")
    return warnings


def generate_content(paras_a, paras_b, title, grade, subject, model=DEFAULT_MODEL, client=None) -> dict:
    import anthropic

    client = client or anthropic.Anthropic()  # reads ANTHROPIC_API_KEY
    prompt = USER_PROMPT.format(
        grade=grade, subject=subject or "General", title=title,
        section_a="\n\n".join(paras_a), section_b="\n\n".join(paras_b) or "(none)",
    )
    last_err = None
    for _ in range(2):  # one retry on malformed JSON
        msg = client.messages.create(model=model, max_tokens=5000, system=SYSTEM_PROMPT,
                                     messages=[{"role": "user", "content": prompt}])
        raw = "".join(b.text for b in msg.content if b.type == "text")
        try:
            content = _parse_json(raw)
            content.setdefault("vocab", [])
            content.setdefault("questions", [])
            content.setdefault("written", {"summary_prompt": "Summarize the main idea of the text below.",
                                           "thesis_prompt": "", "comprehension_prompt": ""})
            content["warnings"] = validate_content(content, "\n".join(paras_a))
            return content
        except (json.JSONDecodeError, KeyError, ValueError) as e:
            last_err = e
    raise RuntimeError(f"Model did not return valid JSON: {last_err}")


# ======================================================================================
# 3. Layout: one neutral list of blocks, rendered to Word or PDF
# ======================================================================================

VOCAB_RUBRIC = [
    ("0", "Student response shows no understanding of the target word. No response or incorrect response."),
    ("1", "Student has partial understanding of the target word. Cannot use the word appropriately in context."),
    ("1.5", "Student has understanding of the target word. Can use it in context but cannot define it accurately."),
    ("2", "Student has full understanding of the target word. Can use it in context and define it accurately."),
]
COMP_RUBRIC = [
    ("0", "Response shows no understanding of the question or text and includes none of the key elements of the suggested answer."),
    ("1", "Response shows limited understanding of the question or text and includes only a minimal part of the suggested answer."),
    ("1.5", "Response shows partial understanding of the question and text and includes some key elements of the suggested answer."),
    ("2", "Response shows clear understanding of the question and text and includes all key elements of the suggested answer."),
]


def decoding_thresholds(n_words: int) -> dict:
    """Independent >=95% accuracy, Instructional 90-95%, Frustration <90%."""
    ind_max, ins_max = round(n_words * 0.05), round(n_words * 0.10)
    return {"ind": (0, ind_max), "ins": (ind_max + 1, ins_max), "fru": ins_max + 1}


def R(text, **o):  # a text run: bold/italic/underline/color/size
    return (text, o)


def P(content, size=11, bold=False, italic=False, color=None, align="left", before=0, after=6, line=1.15, keep=False):
    runs = [R(content)] if isinstance(content, str) else list(content)
    return {"t": "p", "runs": runs, "size": size, "bold": bold, "italic": italic, "color": color,
            "align": align, "before": before, "after": after, "line": line, "keep": keep}


def LINES(n, gap=20):
    return {"t": "lines", "n": n, "gap": gap}


def TABLE(rows, widths, size=9.5, header=True):
    return {"t": "table", "rows": rows, "widths": widths, "size": size, "header": header}


PB = {"t": "pb"}


def build_blocks(title, grade, subject, paras_a, paras_b, content, mode="teacher"):
    """Return (blocks, footer_text). mode: 'teacher' or 'student'."""
    S = mode == "student"
    C = content
    short = C.get("short_title") or title
    allp = paras_a + paras_b
    n_words = sum(word_count(p) for p in paras_a)
    th = decoding_thresholds(n_words)
    B: list = []

    def tag(t):
        if not S:
            B.append(P(t, bold=True, size=8, color=GRAY, align="right", after=0))

    def labeled(label, text, size=11):
        B.append(P([R(label, bold=True), R(text)], size=size))

    box = "[   ]"

    def running_record():
        B.append(PB)
        tag("TEACHER PAGE")
        if S:
            B.append(P("TEACHER COPY \u2013 running record. Remove this page before giving the packet to the student.",
                       bold=True, size=9, color=RED, align="right", after=0))
        B.append(P(f"{title} \u2013 Teacher Copy for Running Record ({n_words} words)", bold=True, size=12, after=4))
        B.append(P("Key: line through a missed word and write what the student said. Mark omissions, insertions, "
                   "repetitions (R) and self-corrections (SC). Read an unknown word aloud, circle it, and note the "
                   "student's definition. [100] markers show cumulative word count.", italic=True, size=9, color=GRAY))
        count, nxt = 0, 100
        for i, para in enumerate(paras_a, 1):
            runs, buf = [R(f"({i}) ", bold=True)], []
            for w in para.split():
                buf.append(w)
                count += 1
                if count == nxt:
                    runs += [R(" ".join(buf) + " "), R(f"[{count}] ", size=8, color=GRAY, bold=True)]
                    buf, nxt = [], nxt + 100
            if buf:
                runs.append(R(" ".join(buf)))
            B.append(P(runs, line=2.0, after=4))

    # ---------------- first page ----------------
    if not S:
        tag("TEACHER PAGE")
        B.append(P("DWIGHT SCHOOL", bold=True, size=20, color=NAVY, align="center", after=2))
        B.append(P(f"QUEST ATL Skill Assessment Gr {grade} {subject + ': ' if subject else ''}{title} #1",
                   bold=True, size=12, align="center", after=10))
        labeled("Primary objective: ", "Develop students' active-reading skills with immediate feedback while reinforcing the concepts and content necessary for class.")
        labeled("Secondary objective: ", "Collect reading-level data for Department of Education (DOE) processes and Quest reports.")
        B.append(P("The main goal of this activity is to gather useful information about the student's literacy skills that will inform your work in Quest and with the classroom teacher while providing needed information for the DOE. Much of this information will also be useful in writing your T1 Quest Report."))
        B.append(P("Prior to Starting with the Student", bold=True, before=8, after=2))
        pn = (C.get("proper_nouns") or [])[:4]
        B.append(P("Print the student reading and student written-response pages single-sided. Keep a separate teacher copy to track the student's oral decoding. Record substitutions, omissions, insertions, repetitions, and self-corrections directly on the teacher copy."
                   + (f" Note difficulty with proper nouns and technical terms (for example, {', '.join(pn)}) separately; do not count these as decoding errors." if pn else "")))
        B.append(P("After Completing the Assessment", bold=True, before=8, after=2))
        B.append(P(f"Create a PDF with the running-record decoding pages and the student's answers to the comprehension questions. Title it Last Name, First Name {short} Skill Assessment and add it to the student's Work Samples folder."))
        B.append(P("Assessment Components", bold=True, before=8, after=2))
        B.append(P("This assessment gathers information in decoding, vocabulary comprehension, explicit and implicit reading comprehension, summarizing, and analytical/thesis writing."))
        if paras_b:
            B.append(P(f"Note: the running record, vocabulary and comprehension questions cover the first {n_words} words. The remaining text is read silently and supports the written responses.", italic=True))
        B.append(PB)
    else:
        B.append(P("DWIGHT SCHOOL", bold=True, size=20, color=NAVY, align="center", after=2))
        B.append(P(f"ATL Activity: {title}", bold=True, size=12, align="center", after=10))
        B.append(P("Name: ______________________________   Date: ______________   Class: ______________", after=10))
        B.append(P("Directions: Read the text carefully. Then complete each section on the pages that follow. Write in complete sentences and use evidence from the text.", italic=True, size=10, color=GRAY))

    # ---------------- student reading ----------------
    tag("STUDENT PAGE")
    B.append(P(title, bold=True, size=14, after=8))
    for i, p in enumerate(allp, 1):
        B.append(P(f"({i}) {p}", size=13, line=1.5, after=8))

    if not S:
        running_record()

    # ---------------- vocabulary ----------------
    def vocab_runs(i, v):
        s, tgt = v.get("sentence", ""), v.get("target", "")
        k = s.lower().find(tgt.lower()) if tgt else -1
        runs = [R(f"{i}. \u201c\u2026")]
        if k >= 0:
            runs += [R(s[:k]), R(s[k:k + len(tgt)], bold=True, underline=True, color=None if S else RED), R(s[k + len(tgt):])]
        else:
            runs.append(R(s))
        runs.append(R("\u2026\u201d"))
        return runs

    B.append(PB)
    if S:
        B.append(P("Vocabulary", bold=True, size=13, after=2))
        B.append(P("Explain what the bold, underlined word or phrase means in your own words. Then use it in a sentence of your own."))
        for i, v in enumerate(C["vocab"], 1):
            B.append(P(vocab_runs(i, v), after=2, keep=True))
            B.append(P("Meaning and sentence:", size=9, color=GRAY, after=0))
            B.append(LINES(3, 18))
            B.append(P("", after=4))
    else:
        tag("TEACHER PAGE")
        B.append(P("Vocabulary Comprehension", bold=True, size=13, after=2))
        B.append(P("Record student responses as accurately as possible. Draw the student's attention to the lines listed below. Ask them to explain the meaning, emphasizing the bold, underlined red word. Award points based on the following:"))
        B.append(TABLE([[a for a, _ in VOCAB_RUBRIC], [b for _, b in VOCAB_RUBRIC]], [0.25] * 4, 8))
        if C.get("preview_words"):
            B.append(P("Additional preview words (unscored): " + ", ".join(C["preview_words"]) + ". Check these informally if time allows and note the student's responses on the teacher copy.", size=9, italic=True))
        for i, v in enumerate(C["vocab"], 1):
            B.append(P(vocab_runs(i, v), after=2, keep=True))
            labeled("Teacher target: ", v.get("teacher_target", ""), 10)
            B.append(P("Student response:", size=9, color=GRAY, after=0))
            B.append(LINES(2, 16))
            B.append(P("Score: ____ / 2", size=9, align="right", after=6))
        B.append(P("Total Score: __________ / 10", bold=True, before=4))

    # ---------------- comprehension ----------------
    B.append(PB)
    if S:
        B.append(P("Reading Comprehension", bold=True, size=13, after=2))
        B.append(P("Answer each question in complete sentences. Use details from the text."))
        for i, q in enumerate(C["questions"], 1):
            B.append(P(f"{i}. {q.get('question', '')}", bold=True, after=2, before=8, keep=True))
            B.append(LINES(4, 18))
    else:
        tag("TEACHER PAGE")
        B.append(P("Reading Comprehension", bold=True, size=13, after=2))
        B.append(P("Record student responses as accurately as possible. Award points based on the following:"))
        B.append(TABLE([[a for a, _ in COMP_RUBRIC], [b for _, b in COMP_RUBRIC]], [0.25] * 4, 8))
        for i, q in enumerate(C["questions"], 1):
            B.append(P(f"{i}. ({q.get('type', '')}) {q.get('question', '')}", bold=True, after=2, before=6, keep=True))
            labeled("Full credit answers should include: ", q.get("full_credit", ""), 10)
            if q.get("distinction"):
                B.append(P(q["distinction"], size=10, italic=True))
            B.append(P("Student response:", size=9, color=GRAY, after=0))
            B.append(LINES(3, 18))
            B.append(P("Score: ____ / 2", size=9, align="right", after=4))
        B.append(P("Calculate the comprehension score: _______ / 10", bold=True, before=4))

    # ---------------- written responses (student pages) ----------------
    W = C["written"]
    B.append(PB)
    tag("STUDENT PAGE")
    B.append(P("Written Responses", bold=True, size=13, after=2))
    B.append(P("Student Name: ________________________________________", after=10))
    B.append(P(W.get("summary_prompt", "Summarize the main idea of the text below."), bold=True, after=2))
    B.append(LINES(7))
    B.append(P("Thesis Practice: " + W.get("thesis_prompt", ""), bold=True, before=14, after=2))
    B.append(LINES(6))
    B.append(PB)
    tag("STUDENT PAGE")
    B.append(P("Written Comprehension: " + W.get("comprehension_prompt", ""), bold=True, after=2))
    B.append(LINES(10))

    if S:
        running_record()
    else:
        # ---------------- writing observations ----------------
        B.append(PB)
        tag("TEACHER PAGE")
        B.append(P("Written Response Observations", bold=True, size=13, after=2))
        B.append(P("Use this page to note the student's writing needs. Check what you see:", after=4))
        B.append(TABLE([
            ["Skill", "Summary", "Thesis practice", "Written comprehension"],
            ["Identifies main idea / claim", box, box, box],
            ["Uses own words (not copied)", box, box, box],
            ["Takes a clear, arguable position", "n/a", box, box],
            ["Supports with specific text evidence", "n/a", box, box],
            ["Organization / logical sequence", box, box, box],
            ["Sentence structure and complete sentences", box, box, box],
            ["Spelling, punctuation, conventions", box, box, box],
            ["Stamina / amount produced", box, box, box],
        ], [0.42, 0.15, 0.2, 0.23]))
        B.append(P("Notes / next steps for Quest and the classroom teacher:", bold=True, before=6, after=2))
        B.append(LINES(8))

        # ---------------- final score report ----------------
        B.append(PB)
        tag("TEACHER PAGE")
        B.append(P("Final Score Report", bold=True, size=13, after=2))
        B.append(P("Complete this after the student has read and orally responded to the vocabulary and comprehension questions. The combined scores on decoding and comprehension allow you to assess whether the student is at the independent, instructional, or frustration level with a text at this level."))
        B.append(TABLE([["", "Student Accuracy Rates"], ["Decoding Score (errors / accuracy %)", ""],
                        ["Vocabulary Comprehension (/10)", ""], ["Comprehension Questions (/10)", ""],
                        ["Total Comprehension (/20)", ""],
                        ["Text Level (Independent / Instructional / Frustration)", ""]], [0.57, 0.43]))
        if C.get("lexile_estimate"):
            B.append(P(f"Estimated Lexile range: {C['lexile_estimate']} (approximate - verify with a Lexile analyzer). Note if your student has significant trouble decoding this text.", size=10))
        B.append(P(f"Count the number of decoding errors in the running-record section ({n_words} words): ____________", bold=True, before=6))
        B.append(TABLE([["Level", "Accuracy", f"Errors ({n_words} words)"],
                        ["Independent", "95% or higher", f"{th['ind'][0]}\u2013{th['ind'][1]}"],
                        ["Instructional", "90\u201395%", f"{th['ins'][0]}\u2013{th['ins'][1]}"],
                        ["Frustration", "below 90%", f"{th['fru']}+"]], [0.33, 0.33, 0.34]))
        B.append(P(f"Accuracy % = ({n_words} \u2212 errors) \u00f7 {n_words} \u00d7 100", size=10, italic=True))
        labeled("Independent: ", "Text is highly accessible for the reader, decoding with 95% or higher accuracy and 90% or higher comprehension (total 18\u201320 of 20).", 10)
        labeled("Instructional: ", "Text is challenging but manageable, decoding with 90\u201395% accuracy and 75% or higher comprehension (total 15\u201317 of 20).", 10)
        labeled("Frustration: ", "Text is too hard for the reader, decoding with below 90% accuracy and/or below 75% comprehension (total below 15 of 20).", 10)
        B.append(P("Adjust cut-offs to match your school's protocol if they differ.", size=9, italic=True, color=GRAY))

    footer = f"Quest ATL {'Activity' if S else 'Skill Assessment'} \u2013 {short}"
    return B, footer


# ======================================================================================
# 4. Renderers
# ======================================================================================


def render_docx(blocks, footer_text, logo=None) -> bytes:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import (WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER)
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(1)
    sec.top_margin = sec.bottom_margin = Inches(0.9)
    st = d.styles["Normal"]
    st.font.name = "Arial"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    st.font.size = Pt(11)

    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = fp.add_run(footer_text + "   ")
    fr.font.size, fr.font.color.rgb = Pt(9), RGBColor.from_string(GRAY)
    run = fp.add_run()
    for kind in ("begin", "instr", "end"):
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = " PAGE "
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        run._r.append(el)
    run.font.size, run.font.color.rgb = Pt(9), RGBColor.from_string(GRAY)

    align = {"left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT, "center": WD_ALIGN_PARAGRAPH.CENTER}
    pending_break = False

    def shade(cell, fill):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        tcPr.append(shd)

    for b in blocks:
        t = b["t"]
        if t == "pb":
            pending_break = True
        elif t == "p":
            p = d.add_paragraph()
            if pending_break:
                p.paragraph_format.page_break_before = True
                pending_break = False
            pf = p.paragraph_format
            pf.space_before, pf.space_after = Pt(b["before"]), Pt(b["after"])
            pf.line_spacing = b["line"]
            pf.keep_with_next = b["keep"] or None
            p.alignment = align[b["align"]]
            for text, o in b["runs"]:
                if not text:
                    continue
                r = p.add_run(text)
                r.bold = o.get("bold", b["bold"]) or None
                r.italic = o.get("italic", b["italic"]) or None
                r.underline = o.get("underline") or None
                r.font.size = Pt(o.get("size", b["size"]))
                col = o.get("color", b["color"])
                if col:
                    r.font.color.rgb = RGBColor.from_string(col)
        elif t == "lines":
            for _ in range(b["n"]):
                p = d.add_paragraph()
                pf = p.paragraph_format
                pf.space_after = Pt(0)
                pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
                pf.line_spacing = Pt(b["gap"])
                pf.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.LINES)
                p.add_run("\t")
        elif t == "table":
            rows, widths = b["rows"], b["widths"]
            tbl = d.add_table(rows=len(rows), cols=len(rows[0]))
            tbl.style = "Table Grid"
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            for i, row in enumerate(rows):
                for j, val in enumerate(row):
                    c = tbl.cell(i, j)
                    c.width = Inches(6.5 * widths[j])
                    c.text = ""
                    r = c.paragraphs[0].add_run(str(val))
                    r.font.size = Pt(b["size"])
                    c.paragraphs[0].paragraph_format.space_after = Pt(2)
                    if b["header"] and i == 0:
                        r.bold = True
                        shade(c, "E8EAF2")
            d.add_paragraph().paragraph_format.space_after = Pt(2)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def render_pdf(blocks, footer_text) -> bytes:
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import HRFlowable, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    def safe(s: str) -> str:
        # built-in PDF fonts cover Western European text; swap anything else for a close match
        s = s.replace("\u2212", "-").replace("\u2265", ">=").replace("\u2264", "<=")
        return s.encode("cp1252", "replace").decode("cp1252")

    al = {"left": TA_LEFT, "right": TA_RIGHT, "center": TA_CENTER}
    story = []
    W = LETTER[0] - 144

    def markup(runs, b):
        out = []
        for text, o in runs:
            if not text:
                continue
            s = escape(safe(text))
            size = o.get("size", b["size"])
            if o.get("bold", b["bold"]):
                s = f"<b>{s}</b>"
            if o.get("italic", b["italic"]):
                s = f"<i>{s}</i>"
            if o.get("underline"):
                s = f"<u>{s}</u>"
            col = o.get("color", b["color"])
            attrs = f' size="{size}"' + (f' color="#{col}"' if col else "")
            out.append(f"<font{attrs}>{s}</font>")
        return "".join(out) or "&nbsp;"

    for b in blocks:
        t = b["t"]
        if t == "pb":
            story.append(PageBreak())
        elif t == "p":
            lead = b["size"] * (1.2 if b["line"] <= 1.15 else b["line"] * 1.2)
            style = ParagraphStyle("x", fontName="Helvetica", fontSize=b["size"], leading=lead, alignment=al[b["align"]],
                                   spaceBefore=b["before"], spaceAfter=b["after"], keepWithNext=1 if b["keep"] else 0)
            story.append(Paragraph(markup(b["runs"], b), style))
        elif t == "lines":
            for _ in range(b["n"]):
                story.append(HRFlowable(width="100%", thickness=0.6, color=colors.HexColor("#" + GRAY),
                                        spaceBefore=b["gap"], spaceAfter=0))
            story.append(Spacer(1, 4))
        elif t == "table":
            cell = ParagraphStyle("c", fontName="Helvetica", fontSize=b["size"], leading=b["size"] * 1.25)
            head = ParagraphStyle("h", parent=cell, fontName="Helvetica-Bold")
            data = [[Paragraph(escape(safe(str(v))), head if (b["header"] and i == 0) else cell) for v in row]
                    for i, row in enumerate(b["rows"])]
            tbl = Table(data, colWidths=[W * w for w in b["widths"]])
            style = [("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#999999")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                     ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]
            if b["header"]:
                style.append(("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EAF2")))
            tbl.setStyle(TableStyle(style))
            story += [tbl, Spacer(1, 8)]

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 9)
        canvas.setFillColor(colors.HexColor("#" + GRAY))
        canvas.drawRightString(LETTER[0] - 72, 36, f"{safe(footer_text)}   {doc.page}")
        canvas.restoreState()

    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=LETTER, leftMargin=72, rightMargin=72, topMargin=62, bottomMargin=62,
                      title=safe(footer_text)).build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


def render(title, grade, subject, paras_a, paras_b, content, mode="teacher", fmt="pdf") -> bytes:
    blocks, footer = build_blocks(title, grade, subject, paras_a, paras_b, content, mode)
    return render_pdf(blocks, footer) if fmt == "pdf" else render_docx(blocks, footer)


# ======================================================================================
# 5. CLI
# ======================================================================================


def main():
    ap = argparse.ArgumentParser(description="Generate an ATL activity (.docx or .pdf) from a class text.")
    ap.add_argument("file", help=".pdf, .docx, .txt or .md containing the class text")
    ap.add_argument("--title", required=True)
    ap.add_argument("--grade", required=True, help='e.g. 9, 10, 11, "11/12"')
    ap.add_argument("--subject", default="")
    ap.add_argument("--decode-words", type=int, default=0, help="Running record covers about this many words (0 = whole text).")
    ap.add_argument("--mode", choices=["teacher", "student"], default="teacher")
    ap.add_argument("--format", choices=["pdf", "docx"], default="pdf")
    ap.add_argument("--out", default=None)
    ap.add_argument("--model", default=DEFAULT_MODEL)
    ap.add_argument("--dump-json", default=None, help="Save the generated content to this JSON file")
    ap.add_argument("--content-json", default=None, help="Skip the API; build from this (edited) JSON file")
    a = ap.parse_args()

    text = read_file(a.file)
    paras = to_paragraphs(text)
    pa, pb = split_for_decoding(paras, a.decode_words)
    if a.content_json:
        content = json.loads(Path(a.content_json).read_text())
        content.setdefault("warnings", validate_content(content, "\n".join(pa)))
    else:
        content = generate_content(pa, pb, a.title, a.grade, a.subject, a.model)
    if a.dump_json:
        Path(a.dump_json).write_text(json.dumps(content, indent=2, ensure_ascii=False))
    for w in content.get("warnings", []):
        print("Warning:", w, file=sys.stderr)
    out = a.out or f"ATL {a.mode} {a.title}.{a.format}"
    Path(out).write_bytes(render(a.title, a.grade, a.subject, pa, pb, content, a.mode, a.format))
    print("Created", out)


if __name__ == "__main__":
    main()
